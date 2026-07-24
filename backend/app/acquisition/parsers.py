from __future__ import annotations

import csv
import io
import json
from datetime import datetime, timezone
from email.utils import parsedate_to_datetime
from typing import Any

import feedparser
from bs4 import BeautifulSoup, Tag

from app.acquisition.contracts import ParseResult, ParsedRecord


def _naive_utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        return value
    return value.astimezone(timezone.utc).replace(tzinfo=None)


def _parse_date(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return _naive_utc(datetime.fromisoformat(value.replace("Z", "+00:00")))
    except ValueError:
        try:
            return _naive_utc(parsedate_to_datetime(value))
        except (TypeError, ValueError, OverflowError):
            return None


def _css_path(node: Tag) -> str:
    if node.get("id"):
        return f"#{node['id']}"
    parts: list[str] = []
    current: Tag | None = node
    while current is not None and current.name not in {"[document]", "html"}:
        part = current.name
        siblings = [
            sibling
            for sibling in current.parent.find_all(current.name, recursive=False)
        ] if current.parent else []
        if len(siblings) > 1:
            part += f":nth-of-type({siblings.index(current) + 1})"
        parts.append(part)
        current = current.parent if isinstance(current.parent, Tag) else None
    return " > ".join(reversed(parts))


def parse_html(content: bytes, media_type: str = "text/html") -> ParseResult:
    soup = BeautifulSoup(content, "lxml")
    title = soup.title.get_text(" ", strip=True) if soup.title else None
    published = None
    for key, value in (
        ("property", "article:published_time"),
        ("name", "date"),
        ("name", "pubdate"),
        ("itemprop", "datePublished"),
    ):
        meta = soup.find("meta", attrs={key: value})
        if meta and meta.get("content"):
            published = _parse_date(str(meta["content"]))
            if published:
                break

    records: list[ParsedRecord] = []
    for node in soup.select("h1, h2, h3, p, li, th, td"):
        text = node.get_text(" ", strip=True)
        if text:
            records.append(
                ParsedRecord(
                    locator_type="css_selector",
                    locator={"selector": _css_path(node)},
                    text=text,
                    fields={"tag": node.name},
                )
            )
    return ParseResult(media_type, title, published, tuple(records))


def parse_rss(content: bytes, media_type: str = "application/rss+xml") -> ParseResult:
    feed = feedparser.parse(content)
    records: list[ParsedRecord] = []
    for index, entry in enumerate(feed.entries, start=1):
        published = None
        if getattr(entry, "published_parsed", None):
            published = datetime(*entry.published_parsed[:6])
        text = " ".join(
            part
            for part in (
                entry.get("title", ""),
                entry.get("summary", ""),
            )
            if part
        )
        records.append(
            ParsedRecord(
                locator_type="xpath",
                locator={"xpath": f"/rss/channel/item[{index}]"},
                text=BeautifulSoup(text, "lxml").get_text(" ", strip=True),
                fields={
                    "title": entry.get("title"),
                    "url": entry.get("link"),
                    "published_at": published.isoformat() if published else None,
                },
            )
        )
    feed_title = feed.feed.get("title") if getattr(feed, "feed", None) else None
    return ParseResult(media_type, feed_title, None, tuple(records))


def parse_csv(content: bytes, media_type: str = "text/csv") -> ParseResult:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    records = tuple(
        ParsedRecord(
            locator_type="table_cell",
            locator={"sheet": "__csv__", "cell": f"A{row_number}"},
            text=json.dumps(row, ensure_ascii=False, sort_keys=True),
            fields=dict(row),
        )
        for row_number, row in enumerate(reader, start=2)
    )
    return ParseResult(
        media_type,
        None,
        None,
        records,
        {"columns": reader.fieldnames or []},
    )


def parse_xlsx(
    content: bytes,
    media_type: str = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
) -> ParseResult:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    records: list[ParsedRecord] = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            continue
        headers = [
            str(value) if value is not None else f"column_{index}"
            for index, value in enumerate(header_row, start=1)
        ]
        for row_number, values in enumerate(rows, start=2):
            fields = {
                header: value
                for header, value in zip(headers, values)
                if value is not None
            }
            if not fields:
                continue
            records.append(
                ParsedRecord(
                    locator_type="table_cell",
                    locator={"sheet": sheet.title, "cell": f"A{row_number}"},
                    text=json.dumps(fields, ensure_ascii=False, sort_keys=True, default=str),
                    fields=fields,
                )
            )
    return ParseResult(media_type, None, None, tuple(records))


def parse_pdf(content: bytes, media_type: str = "application/pdf") -> ParseResult:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    records = tuple(
        ParsedRecord(
            locator_type="pdf_page",
            locator={"page": page_number},
            text=page.extract_text() or "",
            fields={},
        )
        for page_number, page in enumerate(reader.pages, start=1)
    )
    title = None
    if reader.metadata:
        title = reader.metadata.title
    return ParseResult(
        media_type,
        title,
        None,
        records,
        {"page_count": len(reader.pages)},
    )


def parse_json(
    content: bytes,
    media_type: str = "application/json",
    config: dict[str, Any] | None = None,
) -> ParseResult:
    value: Any = json.loads(content)
    config = config or {}
    record_path = config.get("record_path", "")
    selected = value
    if record_path:
        if not isinstance(record_path, str) or not record_path.startswith("/"):
            raise ValueError("JSON record_path must be a JSON Pointer")
        for token in record_path.lstrip("/").split("/"):
            token = token.replace("~1", "/").replace("~0", "~")
            if isinstance(selected, list):
                selected = selected[int(token)]
            elif isinstance(selected, dict):
                selected = selected[token]
            else:
                raise ValueError(f"JSON record_path cannot traverse {token!r}")
    items = selected if isinstance(selected, list) else [selected]
    pointer_prefix = record_path.rstrip("/")
    records = tuple(
        ParsedRecord(
            locator_type="json_pointer",
            locator={
                "pointer": (
                    f"{pointer_prefix}/{index}"
                    if isinstance(selected, list)
                    else pointer_prefix
                )
            },
            text=json.dumps(item, ensure_ascii=False, sort_keys=True),
            fields=item if isinstance(item, dict) else {"value": item},
        )
        for index, item in enumerate(items)
    )
    return ParseResult(
        media_type,
        None,
        None,
        records,
        {
            "record_path": record_path,
            "record_count": len(records),
        },
    )


PARSERS = {
    "html": parse_html,
    "rss": parse_rss,
    "csv": parse_csv,
    "xlsx": parse_xlsx,
    "pdf": parse_pdf,
    "api": parse_json,
}


def parse_content(
    kind: str,
    content: bytes,
    media_type: str,
    config: dict[str, Any] | None = None,
) -> ParseResult:
    try:
        parser = PARSERS[kind]
    except KeyError as exc:
        raise ValueError(f"unsupported source kind: {kind}") from exc
    if kind == "api":
        return parser(content, media_type, config)
    return parser(content, media_type)
